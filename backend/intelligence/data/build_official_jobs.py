"""Build official_jobs.json from the raw BLS OEWS and O*NET files.

official_jobs.json holds the official (BLS / O*NET) tech jobs the seed command loads.

Sources (kept outside the repo, in the directory passed as --raw-dir):
- BLS OEWS May 2025 national estimates: oesm25nat/national_M2025_dl.xlsx
- O*NET 31.0 database, Excel version: db_31_0_excel/

O*NET data is licensed CC BY 4.0: "This page includes information from the O*NET 31.0
Database by the U.S. Department of Labor, Employment and Training Administration
(USDOL/ETA). Used under the CC BY 4.0 license."

Run from backend/ (openpyxl is pulled in for this run only, not added to pyproject.toml):

    uv run --with openpyxl python intelligence/data/build_official_jobs.py --raw-dir <dir>
"""

import argparse
import json
import logging
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

logger = logging.getLogger('build_official_jobs')

BLS_FILE = 'oesm25nat/national_M2025_dl.xlsx'
BLS_SHEET = 'national_M2025_dl'
ONET_DIR = 'db_31_0_excel'
# O*NET 31 split the old Skills file: Essential = basic skills (2.A), Transferable =
# cross-functional skills (2.B). Together they're the 35 general skills. Each maps to
# its JobSkill.skill_type.
ONET_SKILL_FILES = {
    'Essential Skills.xlsx': 'essential',
    'Transferable Skills.xlsx': 'transferable',
}
SKILL_TYPES = {'essential', 'transferable', 'software'}

# Every detailed occupation in SOC major group 15, plus these tech-adjacent codes.
TECH_MAJOR_GROUP = '15'
EXTRA_CODES = {'11-3021', '17-2061', '17-2071', '17-2072', '25-1021'}

GENERAL_SKILL_COUNT = 10
SOC_CODE_RE = re.compile(r'^\d{2}-\d{4}$')


def read_rows(path, sheet=None):
    """Yield each row of an XLSX sheet as a dict keyed by the header row."""
    workbook = openpyxl.load_workbook(path, read_only=True)
    worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows)
    for row in rows:
        yield dict(zip(header, row, strict=True))
    workbook.close()


def read_top_code(path):
    """The annual wage BLS reports as '#' (at or above the top-code), parsed from the
    notes on the Field Descriptions sheet."""
    workbook = openpyxl.load_workbook(path, read_only=True)
    for row in workbook['Field Descriptions'].iter_rows(values_only=True):
        text = ' '.join(str(cell) for cell in row if cell is not None)
        if text.lstrip().startswith('#'):
            match = re.search(r'\$([\d,]+) per year', text)
            if match:
                workbook.close()
                return int(match.group(1).replace(',', ''))
    workbook.close()
    raise ValueError(f'Could not find the "#" top-code note in {path}')


def parse_wage(code, title, value, top_code):
    if value == '#':
        logger.info('%s %s: median wage "#" -> top-code %d', code, title, top_code)
        return top_code
    if value in ('*', None, ''):
        logger.info('%s %s: median wage %r -> None', code, title, value)
        return None
    return int(float(value))


def parse_employment(code, title, value):
    if value in ('**', None, ''):
        logger.info('%s %s: employment %r -> None', code, title, value)
        return None
    return int(float(value))


def load_bls(raw_dir):
    """The candidate occupations: detailed rows only (drops BLS's total, major, minor
    and broad summary rows), in the tech major group or EXTRA_CODES."""
    path = raw_dir / BLS_FILE
    top_code = read_top_code(path)
    candidates = []
    for row in read_rows(path, BLS_SHEET):
        code = row['OCC_CODE']
        if row['O_GROUP'] != 'detailed':
            continue
        if not (code.startswith(f'{TECH_MAJOR_GROUP}-') or code in EXTRA_CODES):
            continue
        title = row['OCC_TITLE']
        candidates.append(
            {
                'soc_code': code,
                'title': title,
                'employment': parse_employment(code, title, row['TOT_EMP']),
                'median_annual_wage': parse_wage(code, title, row['A_MEDIAN'], top_code),
            }
        )
    candidates.sort(key=lambda c: (-(c['employment'] or 0), c['soc_code']))
    return candidates


def load_onet(raw_dir):
    """Descriptions, general skills ({name: (importance on the IM scale, skill type)})
    and In Demand software skill names, keyed by O*NET-SOC code (e.g. 15-1252.00)."""
    onet_dir = raw_dir / ONET_DIR
    descriptions = {
        row['O*NET-SOC Code']: row['Description']
        for row in read_rows(onet_dir / 'Occupation Data.xlsx')
    }

    general = defaultdict(dict)
    for filename, skill_type in ONET_SKILL_FILES.items():
        for row in read_rows(onet_dir / filename):
            if row['Scale ID'] != 'IM':
                continue
            importance = float(row['Data Value'])
            if not 1 <= importance <= 5:
                raise ValueError(
                    f'{filename}: {row["O*NET-SOC Code"]} {row["Element Name"]} '
                    f'importance {importance} is outside 1-5'
                )
            general[row['O*NET-SOC Code']][row['Element Name']] = (importance, skill_type)

    # Only In Demand items - Hot-only ones are often generic tools (Adobe Acrobat) that
    # crowd out the occupation's real stack. A set dedupes repeated Workplace Examples.
    software = defaultdict(set)
    for row in read_rows(onet_dir / 'Software Skills.xlsx'):
        if row['In Demand'] == 'Y':
            software[row['O*NET-SOC Code']].add(row['Workplace Example'])

    return descriptions, general, software


def build_skills(general, software):
    top_general = sorted(general.items(), key=lambda item: (-item[1][0], item[0]))
    skills = [
        {'name': name, 'importance': round(importance, 2), 'type': skill_type}
        for name, (importance, skill_type) in top_general[:GENERAL_SKILL_COUNT]
    ]
    general_names = {skill['name'] for skill in skills}
    # Every In Demand item, no cap. Sorted only so the output is stable between runs.
    skills += [
        {'name': name, 'importance': None, 'type': 'software'}
        for name in sorted(software - general_names)
    ]
    return skills


def build_jobs(candidates, descriptions, general, software):
    """Join BLS to O*NET on the .00 code only - the .01+ codes are O*NET specialties
    BLS doesn't measure separately. Returns (jobs, skipped)."""
    jobs, skipped = [], []
    found = {c['soc_code'] for c in candidates}
    for code in sorted(EXTRA_CODES - found):
        logger.warning('skip %s: not a detailed occupation in BLS', code)
        skipped.append(code)

    for rank, candidate in enumerate(candidates, start=1):
        code, title = candidate['soc_code'], candidate['title']
        onet_code = f'{code}.00'
        if onet_code not in descriptions:
            reason = f'no O*NET {onet_code}'
        elif onet_code not in general:
            reason = f'no skill ratings for O*NET {onet_code}'
        else:
            reason = None
        if reason:
            logger.warning('skip #%d %s %s: %s', rank, code, title, reason)
            skipped.append(code)
            continue
        jobs.append(
            {
                'title': title,
                'soc_code': code,
                'description': descriptions[onet_code],
                'area': 'National',
                'median_annual_wage': candidate['median_annual_wage'],
                'employment': candidate['employment'],
                'source': 'official',
                'skills': build_skills(general[onet_code], software.get(onet_code, set())),
            }
        )
    return jobs, skipped


def validate(jobs, skipped, candidates):
    kept = [job['soc_code'] for job in jobs]
    expected = {c['soc_code'] for c in candidates} | EXTRA_CODES
    if set(kept) & set(skipped):
        raise ValueError(f'Codes both kept and skipped: {sorted(set(kept) & set(skipped))}')
    if set(kept) | set(skipped) != expected:
        raise ValueError(
            f'Codes neither kept nor skipped: {sorted(expected - set(kept) - set(skipped))}'
        )
    if len(kept) != len(set(kept)):
        raise ValueError('Duplicate soc_code in output')
    for job in jobs:
        code = job['soc_code']
        if not SOC_CODE_RE.match(code):
            raise ValueError(f'{code}: not a BLS-format SOC code')
        names = [skill['name'] for skill in job['skills']]
        if len(names) != len(set(names)):
            raise ValueError(f'{code}: duplicate skill names')
        types = {skill['type'] for skill in job['skills']}
        if not types <= SKILL_TYPES:
            raise ValueError(f'{code}: unknown skill types {sorted(types - SKILL_TYPES)}')
        general = [s for s in job['skills'] if s['type'] != 'software']
        if len(general) != GENERAL_SKILL_COUNT:
            raise ValueError(
                f'{code}: {len(general)} general skills, expected {GENERAL_SKILL_COUNT}'
            )
        if not all(s['importance'] is not None and 1 <= s['importance'] <= 5 for s in general):
            raise ValueError(f'{code}: general skill importance missing or outside 1-5')
        if any(s['importance'] is not None for s in job['skills'] if s['type'] == 'software'):
            raise ValueError(f'{code}: software skills must have null importance')


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        '--raw-dir',
        type=Path,
        required=True,
        help=f'Directory containing {BLS_FILE} and {ONET_DIR}/.',
    )
    parser.add_argument(
        '--out', type=Path, default=Path(__file__).resolve().parent / 'official_jobs.json'
    )
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format='%(levelname)s %(message)s')

    candidates = load_bls(args.raw_dir)
    descriptions, general, software = load_onet(args.raw_dir)
    jobs, skipped = build_jobs(candidates, descriptions, general, software)
    validate(jobs, skipped, candidates)

    args.out.write_text(json.dumps(jobs, indent=2, ensure_ascii=False) + '\n')
    logger.info(
        'candidates %d, kept %d, skipped %d',
        len({c['soc_code'] for c in candidates} | EXTRA_CODES),
        len(jobs),
        len(skipped),
    )
    software_counts = [
        sum(1 for skill in job['skills'] if skill['type'] == 'software') for job in jobs
    ]
    logger.info(
        'software skills per job: min %d, max %d, avg %.1f',
        min(software_counts),
        max(software_counts),
        sum(software_counts) / len(software_counts),
    )
    logger.info('wrote %s', args.out)


if __name__ == '__main__':
    main()
